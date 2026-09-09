import tkinter as tk
from tkinter import ttk, messagebox
from openpyxl import Workbook, load_workbook
import os


# ---------------------------------------------------------
# Excel File
# ---------------------------------------------------------

FILE_NAME = "student_results.xlsx"

COLUMNS = [
    "Name",
    "Roll No.",
    "Class",
    "Subject 1",
    "Subject 2",
    "Subject 3",
    "Subject 4",
    "Subject 5",
    "Total Marks",
    "Percentage",
    "Result"
]


# ---------------------------------------------------------
# Create Excel File if it doesn't exist
# ---------------------------------------------------------

def create_excel_file():
    if not os.path.exists(FILE_NAME):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Student Results"

        for col, heading in enumerate(COLUMNS, start=1):
            sheet.cell(row=1, column=col, value=heading)

        workbook.save(FILE_NAME)


# ---------------------------------------------------------
# Clear all input fields
# ---------------------------------------------------------

def clear_fields():
    name_entry.delete(0, tk.END)
    roll_entry.delete(0, tk.END)
    class_entry.delete(0, tk.END)

    for entry in subject_entries:
        entry.delete(0, tk.END)


# ---------------------------------------------------------
# Save Student
# ---------------------------------------------------------

def save_student():

    name = name_entry.get().strip()
    roll_no = roll_entry.get().strip()
    student_class = class_entry.get().strip()

    # Check basic details
    if name == "" or roll_no == "" or student_class == "":
        messagebox.showerror(
            "Error",
            "Please enter Name, Roll No. and Class."
        )
        return

    # Check roll number
    if not roll_no.isdigit():
        messagebox.showerror(
            "Error",
            "Roll No. must contain numbers only."
        )
        return

    # Get subject marks
    marks = []

    for i, entry in enumerate(subject_entries):
        value = entry.get().strip()

        if value == "":
            messagebox.showerror(
                "Error",
                f"Please enter marks for Subject {i + 1}."
            )
            return

        try:
            mark = float(value)

            if mark < 0 or mark > 100:
                messagebox.showerror(
                    "Error",
                    f"Subject {i + 1} marks must be between 0 and 100."
                )
                return

            marks.append(mark)

        except ValueError:
            messagebox.showerror(
                "Error",
                f"Subject {i + 1} marks must be a number."
            )
            return

    # Check duplicate roll number
    workbook = load_workbook(FILE_NAME)
    sheet = workbook.active

    for row in sheet.iter_rows(min_row=2, values_only=True):
        if str(row[1]) == roll_no:
            messagebox.showerror(
                "Error",
                "A student with this Roll No. already exists."
            )
            workbook.close()
            return

    # Calculate total and percentage
    total = sum(marks)
    percentage = total / 5

    # Pass/Fail
    # Student passes only if every subject is 40 or above
    if all(mark >= 40 for mark in marks):
        result = "Pass"
    else:
        result = "Fail"

    # Add record to Excel
    sheet.append([
        name,
        int(roll_no),
        student_class,
        marks[0],
        marks[1],
        marks[2],
        marks[3],
        marks[4],
        total,
        percentage,
        result
    ])

    workbook.save(FILE_NAME)
    workbook.close()

    messagebox.showinfo(
        "Success",
        f"Student record saved successfully!\n\n"
        f"Total Marks: {total}\n"
        f"Percentage: {percentage:.2f}%\n"
        f"Result: {result}"
    )

    clear_fields()


# ---------------------------------------------------------
# Get Result
# ---------------------------------------------------------

def get_result():

    roll_no = search_roll_entry.get().strip()

    if roll_no == "":
        messagebox.showerror(
            "Error",
            "Please enter Roll No."
        )
        return

    if not roll_no.isdigit():
        messagebox.showerror(
            "Error",
            "Roll No. must contain numbers only."
        )
        return

    workbook = load_workbook(FILE_NAME)
    sheet = workbook.active

    found = False

    for row in sheet.iter_rows(min_row=2, values_only=True):

        if str(row[1]) == roll_no:

            found = True

            result_window = tk.Toplevel(root)
            result_window.title("Student Result")
            result_window.geometry("600x400")
            result_window.resizable(False, False)

            title = tk.Label(
                result_window,
                text="STUDENT RESULT",
                font=("Arial", 20, "bold")
            )
            title.pack(pady=20)

            result_frame = tk.Frame(result_window)
            result_frame.pack(pady=10)

            details = [
                ("Name", row[0]),
                ("Roll No.", row[1]),
                ("Class", row[2]),
                ("Total Marks", row[8]),
                ("Percentage", f"{row[9]:.2f}%"),
                ("Result", row[10])
            ]

            for i, (label, value) in enumerate(details):

                tk.Label(
                    result_frame,
                    text=label + ":",
                    font=("Arial", 13, "bold"),
                    width=18,
                    anchor="w"
                ).grid(row=i, column=0, padx=10, pady=8)

                tk.Label(
                    result_frame,
                    text=value,
                    font=("Arial", 13),
                    width=25,
                    anchor="w"
                ).grid(row=i, column=1, padx=10, pady=8)

            break

    workbook.close()

    if not found:
        messagebox.showerror(
            "Not Found",
            "Student record not found."
        )


# ---------------------------------------------------------
# Show All Results
# ---------------------------------------------------------

def show_all_results():

    workbook = load_workbook(FILE_NAME)
    sheet = workbook.active

    all_window = tk.Toplevel(root)
    all_window.title("All Student Results")
    all_window.geometry("900x500")

    title = tk.Label(
        all_window,
        text="ALL STUDENT RESULTS",
        font=("Arial", 20, "bold")
    )
    title.pack(pady=15)

    # Frame for Treeview
    table_frame = tk.Frame(all_window)
    table_frame.pack(fill="both", expand=True, padx=15, pady=10)

    # Treeview columns
    columns = (
        "Name",
        "Roll No.",
        "Class",
        "Total Marks",
        "Percentage",
        "Result"
    )

    tree = ttk.Treeview(
        table_frame,
        columns=columns,
        show="headings"
    )

    # Headings
    for column in columns:
        tree.heading(column, text=column)

    # Column widths
    tree.column("Name", width=150)
    tree.column("Roll No.", width=100)
    tree.column("Class", width=100)
    tree.column("Total Marks", width=120)
    tree.column("Percentage", width=120)
    tree.column("Result", width=100)

    # Add scrollbar
    scrollbar = ttk.Scrollbar(
        table_frame,
        orient="vertical",
        command=tree.yview
    )

    tree.configure(yscrollcommand=scrollbar.set)

    tree.pack(
        side="left",
        fill="both",
        expand=True
    )

    scrollbar.pack(
        side="right",
        fill="y"
    )

    # Read Excel records
    for row in sheet.iter_rows(min_row=2, values_only=True):

        tree.insert(
            "",
            tk.END,
            values=(
                row[0],
                row[1],
                row[2],
                row[8],
                f"{row[9]:.2f}%",
                row[10]
            )
        )

    workbook.close()


# ---------------------------------------------------------
# Exit
# ---------------------------------------------------------

def exit_program():

    answer = messagebox.askyesno(
        "Exit",
        "Do you want to exit the application?"
    )

    if answer:
        root.destroy()


# ---------------------------------------------------------
# Main Window
# ---------------------------------------------------------

root = tk.Tk()

root.title("Student Result Management System")
root.geometry("900x750")
root.resizable(False, False)


# ---------------------------------------------------------
# Main Heading
# ---------------------------------------------------------

title_label = tk.Label(
    root,
    text="STUDENT RESULT MANAGEMENT SYSTEM",
    font=("Arial", 24, "bold")
)

title_label.pack(pady=20)


# ---------------------------------------------------------
# Add Student Frame
# ---------------------------------------------------------

add_frame = tk.LabelFrame(
    root,
    text="Add Student",
    font=("Arial", 15, "bold"),
    padx=20,
    pady=15
)

add_frame.pack(
    padx=30,
    pady=10,
    fill="x"
)


# Name
tk.Label(
    add_frame,
    text="Name:",
    font=("Arial", 12)
).grid(row=0, column=0, padx=10, pady=8, sticky="w")

name_entry = tk.Entry(
    add_frame,
    font=("Arial", 12),
    width=25
)

name_entry.grid(
    row=0,
    column=1,
    padx=10,
    pady=8
)


# Roll Number
tk.Label(
    add_frame,
    text="Roll No.:",
    font=("Arial", 12)
).grid(row=0, column=2, padx=10, pady=8, sticky="w")

roll_entry = tk.Entry(
    add_frame,
    font=("Arial", 12),
    width=25
)

roll_entry.grid(
    row=0,
    column=3,
    padx=10,
    pady=8
)


# Class
tk.Label(
    add_frame,
    text="Class:",
    font=("Arial", 12)
).grid(row=1, column=0, padx=10, pady=8, sticky="w")

class_entry = tk.Entry(
    add_frame,
    font=("Arial", 12),
    width=25
)

class_entry.grid(
    row=1,
    column=1,
    padx=10,
    pady=8
)


# Subject Entries
subject_entries = []

for i in range(5):

    row_number = i + 2

    tk.Label(
        add_frame,
        text=f"Subject {i + 1} Marks:",
        font=("Arial", 12)
    ).grid(
        row=row_number,
        column=0,
        padx=10,
        pady=6,
        sticky="w"
    )

    entry = tk.Entry(
        add_frame,
        font=("Arial", 12),
        width=25
    )

    entry.grid(
        row=row_number,
        column=1,
        padx=10,
        pady=6
    )

    subject_entries.append(entry)


# Save Button
save_button = tk.Button(
    add_frame,
    text="💾 Save",
    font=("Arial", 12, "bold"),
    width=15,
    command=save_student
)

save_button.grid(
    row=7,
    column=0,
    columnspan=2,
    pady=15
)


# Clear Button
clear_button = tk.Button(
    add_frame,
    text="Clear",
    font=("Arial", 12, "bold"),
    width=15,
    command=clear_fields
)

clear_button.grid(
    row=7,
    column=2,
    columnspan=2,
    pady=15
)


# ---------------------------------------------------------
# Get Result Frame
# ---------------------------------------------------------

search_frame = tk.LabelFrame(
    root,
    text="Get Result",
    font=("Arial", 15, "bold"),
    padx=20,
    pady=15
)

search_frame.pack(
    padx=30,
    pady=10,
    fill="x"
)


tk.Label(
    search_frame,
    text="Enter Roll No.:",
    font=("Arial", 12)
).grid(
    row=0,
    column=0,
    padx=10,
    pady=10
)


search_roll_entry = tk.Entry(
    search_frame,
    font=("Arial", 12),
    width=25
)

search_roll_entry.grid(
    row=0,
    column=1,
    padx=10,
    pady=10
)


get_button = tk.Button(
    search_frame,
    text="🔍 Get Result",
    font=("Arial", 12, "bold"),
    width=15,
    command=get_result
)

get_button.grid(
    row=0,
    column=2,
    padx=20,
    pady=10
)


# ---------------------------------------------------------
# Main Action Buttons
# ---------------------------------------------------------

button_frame = tk.Frame(root)

button_frame.pack(pady=20)


show_button = tk.Button(
    button_frame,
    text="📋 Show All Results",
    font=("Arial", 13, "bold"),
    width=20,
    height=2,
    command=show_all_results
)

show_button.grid(
    row=0,
    column=0,
    padx=15
)


exit_button = tk.Button(
    button_frame,
    text="Exit",
    font=("Arial", 13, "bold"),
    width=15,
    height=2,
    command=exit_program
)

exit_button.grid(
    row=0,
    column=1,
    padx=15
)


# ---------------------------------------------------------
# Create Excel file and start application
# ---------------------------------------------------------

create_excel_file()

root.mainloop()